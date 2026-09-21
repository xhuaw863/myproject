# -*- coding: utf-8 -*-
"""医疗服务项目三文档独立完整抽取(不合并、不做跨版本对照)。

按用户要求: 三个文档各自原样完整导入为独立字典, 暂不做来源间对照/合并。
 - std_msi_nat  全国医疗服务项目技术规范(2023年版).xlsx Sheet1: 表头第704行,
   8位字母数字混合项目码(如 CAA01LB1)为数据行, 1-4位纯字母为层级标题(记入 cat_name 路径);
   原生列全保留(中英文名/内涵/必需耗材/可选耗材/低值耗材分档/基本人力消耗及耗时/
   技术难度/风险程度/人力资源消耗相对值/计量单位/说明/特殊情况资源消耗调整系数/
   收费票据分类/会计科目分类/病案首页费用分类)。
 - std_msi_hb   湖北省医疗服务价格项目及医保支付目录(2023版).xlsx sheet"湖北医保物价目录":
   9位数字码=基础项, 9位+字母后缀=子项, 均原样成行; 2/4位数字与"一、"行为层级标题(记入 cat_name)。
 - std_msi_fin  医疗服务项目相关财务归集口径规范.pdf: 827页线框表逐行原样导入
   (2023码/名, 2012码/名, 2001码/名, 收费票据分类, 会计科目分类, 病案首页费用分类),
   层级行(A/AA/…)与续行(2023码为空)保留原样, 仅跳过表头行。
输出: resources/seed/std_msi_nat.tsv / std_msi_hb.tsv / std_msi_fin.tsv (首行表头) + 统计。
"""
import os
import re
import openpyxl
import pdfplumber

BASE = r"D:\study\ybtest\字典标准\国家和湖北省医疗物价规范"
SEED = os.path.join(r"D:\study\ybtest\yb-interface", "src", "main", "resources", "seed")
STATS = r"D:\study\ybtest\_msi_stats.out"

ITEM_RE = re.compile(r"^(?=.*\d)[A-Z0-9]{8}$")
stats = []


def cl(v):
    """清洗单元格: 去换行/制表, 压缩空白。"""
    if v is None:
        return ""
    s = str(v).replace("\r", "").replace("\n", "").replace("\t", " ")
    return re.sub(r"\s+", " ", s).strip()


def cl_en(v):
    """英文名: 换行转空格保留单词间隔。"""
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v).replace("\r", "").replace("\n", " ").replace("\t", " ")).strip()


def write_tsv(name, header, rows):
    path = os.path.join(SEED, name + ".tsv")
    with open(path, "w", encoding="utf-8", newline="") as f:
        f.write("\t".join(header) + "\n")
        for r in rows:
            f.write("\t".join(r) + "\n")
    return len(rows)


# ---------------- 1. 国家 2023 技术规范 xlsx (原样全列) ----------------
def load_national():
    path = os.path.join(BASE, "全国医疗服务项目技术规范(2023年版).xlsx")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Sheet1"]
    rows = []
    path_lv = ["", "", "", ""]
    hdr = -1
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if hdr < 0:
            if row and cl(row[0]) == "项目编码":
                hdr = i
            continue
        code = cl(row[0]) if row and row[0] is not None else ""
        if not code:
            continue
        name = cl(row[2]) if len(row) > 2 else ""

        def g(idx):
            return cl(row[idx]) if len(row) > idx and row[idx] is not None else ""

        if ITEM_RE.match(code):
            if not name:
                # 附录区(一次性医用耗材分类: 器械和器具等)项目行仅有编码无名称, 不纳入字典
                continue
            rows.append([code, name, cl_en(row[5]) if len(row) > 5 else "", g(7), g(16), g(18),
                         g(20), g(21), g(24), g(25), g(26), g(28), g(29), g(30), g(32), g(34),
                         g(35), " > ".join([p for p in path_lv if p])])
        elif len(code) == 1:
            path_lv = [name, "", "", ""]
        elif len(code) == 2:
            path_lv[1] = name
            path_lv[2] = path_lv[3] = ""
        elif len(code) == 3:
            path_lv[2] = name
            path_lv[3] = ""
        elif len(code) == 4 and code.isalpha():
            path_lv[3] = name
        # 其余(5-7位区间行、附录耗材分类文字行等)非项目数据, 忽略
    wb.close()
    header = ["item_code", "item_name", "item_name_en", "item_content", "consumable_req",
              "consumable_opt", "consumable_low", "hr_time", "tech_difficulty", "risk_level",
              "hr_value", "unit", "remark", "adjust_coef", "invoice_class", "acct_class",
              "mr_cost_class", "cat_name"]
    n = write_tsv("std_msi_nat", header, rows)
    stats.append("msi_nat rows=%d" % n)
    return n


# ---------------- 2. 湖北 2023 物价目录 xlsx (基础项+子项原样) ----------------
def load_hubei():
    path = os.path.join(BASE, "湖北省医疗服务价格项目及医保支付目录(2023版).xlsx")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["湖北医保物价目录"]
    rows = []
    lv = ["", "", ""]
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 3:
            continue
        code = cl(row[0]) if row and row[0] is not None else ""
        if not code:
            continue
        name = cl(row[1]) if len(row) > 1 else ""

        def g(idx):
            return cl(row[idx]) if len(row) > idx and row[idx] is not None else ""

        if re.match(r"^[一二三四五六七八九十]+、", code):
            lv = [code, "", ""]
        elif re.match(r"^\d{2}$", code):
            lv[1] = name
            lv[2] = ""
        elif re.match(r"^\d{4}$", code):
            lv[2] = name
        elif re.match(r"^\d{9}([a-zA-Z]+)?$", code):
            rows.append([code, name, g(2), g(3), g(4), g(5), g(6), g(7), g(8),
                         " > ".join([p for p in lv if p])])
    wb.close()
    header = ["item_code", "item_name", "item_content", "excluded", "unit", "pay_cat",
              "item_explain", "remark", "trial", "cat_name"]
    n = write_tsv("std_msi_hb", header, rows)
    stats.append("msi_hb rows=%d" % n)
    return n


# ---------------- 3. 财务归集口径 PDF (表行原样, 含层级行/续行) ----------------
def load_fin():
    path = os.path.join(BASE, "医疗服务项目相关财务归集口径规范.pdf")
    rows = []
    pdf = pdfplumber.open(path)
    start = None
    for i in range(2, 60):
        t = pdf.pages[i].extract_text() or ""
        if ITEM_RE.search(t):
            start = i
            break
    if start is None:
        start = 25
    for pi in range(start, len(pdf.pages)):
        for tb in pdf.pages[pi].extract_tables():
            for row in tb:
                cells = [cl(c) for c in row[:9]] + [""] * (9 - min(len(row), 9))
                if cells[0] == "项目编码" or cells[0].startswith("全国医疗服务"):
                    continue
                if not any(cells):
                    continue
                rows.append(cells)
    pdf.close()
    header = ["code_2023", "name_2023", "code_2012", "name_2012", "code_2001", "name_2001",
              "invoice_class", "acct_class", "mr_cost_class"]
    n = write_tsv("std_msi_fin", header, rows)
    stats.append("msi_fin rows=%d (pages %d..%d)" % (n, start, len(pdf.pages) - 1))
    return n


def main():
    os.makedirs(SEED, exist_ok=True)
    a = load_national()
    b = load_hubei()
    c = load_fin()
    open(STATS, "w", encoding="utf-8").write("\n".join(stats))
    print("done nat=%d hb=%d fin=%d" % (a, b, c))


if __name__ == "__main__":
    main()
